"""Teacher question bank: Excel upload parsing + student bank-first pick.

Teacher Excel upload (teammate UI) stores rows in ``BankQuestion`` /
``teacher_bank_questions``. Student generate calls ``pick_question`` first.

Course matching: teachers should enter the student chapter label (or the
chapter_id) as the Excel ``course`` field. Pick accepts either.
"""

from __future__ import annotations

import hashlib
import io
import json
import random
import uuid
import zipfile
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import or_

from common.models import BankQuestion, TeacherQuestionBankItem, db


HEADERS = [
    "question_id", "question_type", "topic", "question",
    "option_A", "option_B", "option_C", "option_D", "option_E",
    "correct_answer", "explanation",
]
MAX_FILE_BYTES = 5 * 1024 * 1024

VALID_TYPES = {"true_false", "mcq", "matching", "fill_blank", "short_answer"}
STUDENT_TO_BANK_TYPE = {"true_false": "TF", "mcq": "MCQ"}
BANK_TO_STUDENT_TYPE = {"TF": "true_false", "MCQ": "mcq"}


def parse_workbook(content):
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("Maximum file size is 5 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 20 MB.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Invalid Excel workbook. Please use the .xlsx template.") from exc
    try:
        sheet = book.active
        # Iterate actual cells as well: producer dimension metadata is not reliable.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        first = next(rows, ())
        if [str(c.value or "").strip() for c in first] != HEADERS:
            raise ValueError("Column headers must match the downloaded template exactly.")
        questions, errors, seen = [], [], set()
        for number, cells in enumerate(rows, 2):
            if number > 1001:
                raise ValueError("Maximum 1000 question rows per upload.")
            if len(cells) > len(HEADERS):
                raise ValueError("Unexpected columns. Use the downloaded template.")
            if not any(c.value is not None for c in cells):
                continue
            values = [str(c.value).strip() if c.value is not None else "" for c in cells]
            values += [""] * (len(HEADERS) - len(values))
            row = dict(zip(HEADERS, values))
            issues = []
            if any(c.data_type == "f" for c in cells):
                issues.append("Formulas are not accepted; paste values instead.")
            for field, limit in [("question_id", 80), ("topic", 256), ("question", 10000), ("explanation", 10000)]:
                if len(row[field]) > limit:
                    issues.append(f"{field} exceeds {limit} characters.")
            if any(len(row["option_" + k]) > 10000 for k in "ABCDE"):
                issues.append("An option exceeds 10000 characters.")
            if not row["question_id"] or not row["question"]:
                issues.append("question_id and question are required.")
            key = row["question_id"].casefold()
            if key in seen:
                issues.append("Duplicate question_id in this upload (case-insensitive).")
            seen.add(key)
            kind = row["question_type"].upper()
            answer = row["correct_answer"].upper()
            options = {k: row["option_" + k] for k in "ABCDE" if row["option_" + k]}
            if kind == "TF":
                if answer not in ("TRUE", "FALSE"):
                    issues.append("TF answer must be TRUE or FALSE.")
                if options:
                    issues.append("Leave option columns empty for TF.")
            elif kind == "MCQ":
                if len(options) < 2:
                    issues.append("MCQ requires at least two options.")
                if answer not in options:
                    issues.append("MCQ answer must be one existing option letter A–E.")
                if len({v.casefold() for v in options.values()}) != len(options):
                    issues.append("Duplicate option text is not accepted.")
            else:
                issues.append("question_type must be TF or MCQ.")
            if issues:
                errors.append({"row": number, "errors": issues})
            else:
                questions.append(dict(
                    row=number,
                    question_id=row["question_id"],
                    question_type=kind,
                    topic=row["topic"],
                    question=row["question"],
                    options=options,
                    correct_answer=answer,
                    explanation=row["explanation"],
                ))
        if not questions and not errors:
            raise ValueError("The workbook contains no questions.")
        return questions, errors
    finally:
        book.close()


def _loads(value: Any, fallback: Any) -> Any:
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value or "")
    except (TypeError, json.JSONDecodeError):
        return fallback


def _dumps(value: Any) -> str:
    if isinstance(value, str):
        raw = value.strip()
        if raw.startswith("{") or raw.startswith("["):
            return raw
        return json.dumps(value, ensure_ascii=False)
    return json.dumps(value if value is not None else {}, ensure_ascii=False)


def quiz_from_bank_question(row: BankQuestion) -> dict[str, Any] | None:
    """Map a teacher-uploaded BankQuestion into the student quiz payload."""
    question = str(row.question or "").strip()
    if not question:
        return None
    kind = str(row.question_type or "").strip().upper()
    explanation = str(row.explanation or "")
    if kind == "TF":
        answer = str(row.correct_answer or "").strip().upper()
        if answer not in {"TRUE", "FALSE"}:
            return None
        return {
            "type": "true_false",
            "question": question,
            "options": ["True", "False"],
            "correct_index": 0 if answer == "TRUE" else 1,
            "explanation": explanation,
        }
    if kind == "MCQ":
        raw_options = row.options if isinstance(row.options, dict) else {}
        letters = [letter for letter in "ABCDE" if str(raw_options.get(letter) or "").strip()]
        if len(letters) < 2:
            return None
        answer = str(row.correct_answer or "").strip().upper()
        if answer not in letters:
            return None
        return {
            "type": "mcq",
            "question": question,
            "options": [str(raw_options[letter]).strip() for letter in letters],
            "correct_index": letters.index(answer),
            "explanation": explanation,
        }
    return None


def quiz_payload_from_item(item: TeacherQuestionBankItem) -> dict[str, Any] | None:
    """Rebuild a submit-ready quiz dict from the reserved JSON-bank table."""
    payload = _loads(item.options_or_pairs, {})
    if not isinstance(payload, dict):
        payload = {}
    question = str(item.question or "").strip()
    if not question:
        return None
    quiz: dict[str, Any] = {
        "type": str(item.question_type or ""),
        "question": question,
        "explanation": str(item.explanation or ""),
    }
    for key in ("options", "correct_index", "left", "right", "correct_mapping", "answer", "reference", "key_points"):
        if key in payload:
            quiz[key] = payload[key]
    quiz_type = quiz["type"]
    if quiz_type in {"true_false", "mcq"} and not quiz.get("options"):
        return None
    if quiz_type == "matching" and (not quiz.get("left") or not quiz.get("right")):
        return None
    if quiz_type == "fill_blank" and "answer" not in quiz:
        return None
    if quiz_type == "short_answer" and not (quiz.get("reference") or quiz.get("answer")):
        return None
    return quiz


def _course_keys(chapter_id: str, chapter_label: str = "") -> list[str]:
    keys: list[str] = []
    for value in (chapter_id, chapter_label):
        text = str(value or "").strip()
        if text and text not in keys:
            keys.append(text)
    if chapter_id and ":" in chapter_id:
        tail = chapter_id.split(":", 1)[-1].replace("-", " ").strip()
        if tail and tail not in keys:
            keys.append(tail)
    return keys


def _hit_from_bank_question(row: BankQuestion, quiz_type: str) -> dict[str, Any] | None:
    quiz = quiz_from_bank_question(row)
    if quiz is None or quiz.get("type") != quiz_type:
        return None
    fingerprint = f"tbank-{row.question_id}"[:80]
    return {
        "quiz": quiz,
        "topic": str(row.topic or ""),
        "fingerprint": fingerprint,
        "generation_source": str(row.source or "teacher_upload"),
        "focus_text": "",
        "fact_key": "",
        "chunk": {
            "chunk_id": "",
            "knowledge_key": "",
            "source_file": f"teacher_bank:{row.course}",
            "page": None,
            "text": "",
        },
        "citations": [],
        "source_excerpt": "",
    }


def _hit_from_legacy_item(item: TeacherQuestionBankItem, quiz_type: str) -> dict[str, Any] | None:
    quiz = quiz_payload_from_item(item)
    if quiz is None or quiz.get("type") != quiz_type:
        return None
    return {
        "quiz": quiz,
        "topic": str(item.topic or ""),
        "fingerprint": str(item.question_fingerprint or f"bank-{item.id}"),
        "generation_source": "question_bank",
        "focus_text": str(item.source_excerpt or item.fact_summary or ""),
        "fact_key": str(item.fact_key or ""),
        "chunk": {
            "chunk_id": str(item.source_chunk_id or ""),
            "knowledge_key": str(item.knowledge_key or ""),
            "source_file": str(item.source_document or ""),
            "page": item.source_page,
            "text": str(item.source_excerpt or ""),
        },
        "citations": _loads(item.citations_json, []),
        "source_excerpt": str(item.source_excerpt or ""),
    }


def pick_question(
    chapter_id: str,
    quiz_type: str,
    topic: str = "",
    used_fingerprints: list[Any] | None = None,
    chapter_label: str = "",
) -> dict[str, Any] | None:
    """Prefer a teacher-bank question before AI generation.

    Order: Excel-uploaded ``BankQuestion`` rows, then reserved JSON bank items.
    """
    chapter_id = str(chapter_id or "").strip()
    quiz_type = str(quiz_type or "").strip()
    if not chapter_id or quiz_type not in VALID_TYPES:
        return None

    used = {str(item) for item in (used_fingerprints or []) if item}
    topic_norm = str(topic or "").strip().lower()
    hits: list[dict[str, Any]] = []

    bank_type = STUDENT_TO_BANK_TYPE.get(quiz_type)
    if bank_type:
        course_keys = _course_keys(chapter_id, chapter_label)
        query = BankQuestion.query.filter_by(question_type=bank_type)
        if course_keys:
            query = query.filter(or_(*[BankQuestion.course == key for key in course_keys]))
        else:
            query = query.filter(False)
        rows = query.order_by(BankQuestion.created_at.desc()).limit(200).all()
        for row in rows:
            hit = _hit_from_bank_question(row, quiz_type)
            if hit is None:
                continue
            if hit["fingerprint"] in used:
                continue
            row_topic = str(row.topic or "").strip().lower()
            if topic_norm and row_topic != topic_norm:
                continue
            hits.append(hit)

    if not hits:
        legacy_rows = (
            TeacherQuestionBankItem.query.filter_by(
                chapter_id=chapter_id,
                question_type=quiz_type,
                is_active=True,
            )
            .order_by(TeacherQuestionBankItem.created_at.desc())
            .limit(200)
            .all()
        )
        for row in legacy_rows:
            hit = _hit_from_legacy_item(row, quiz_type)
            if hit is None:
                continue
            if hit["fingerprint"] in used:
                continue
            row_topic = str(row.topic or "").strip().lower()
            if topic_norm and row_topic != topic_norm:
                continue
            hits.append(hit)

    if not hits:
        return None
    if topic_norm:
        return hits[0]
    return random.choice(hits)


def _fingerprint_for(row: dict[str, Any]) -> str:
    existing = str(row.get("question_fingerprint") or "").strip()
    if existing:
        return existing[:80]
    basis = "|".join(
        [
            str(row.get("chapter_id") or "").strip(),
            str(row.get("question_type") or "").strip(),
            str(row.get("topic") or "").strip().lower(),
            str(row.get("question") or "").strip(),
        ]
    )
    return hashlib.sha1(basis.encode("utf-8")).hexdigest()[:80]


def upsert_items(rows: list[dict[str, Any]], uploaded_by: str = "") -> int:
    """Insert or update reserved JSON bank items (legacy/API helper).

    Prefer the Excel Question bank UI for teacher uploads. This remains for
    ``POST /question_bank/items`` automation.
    """
    if not isinstance(rows, list):
        raise ValueError("rows must be a list of question objects")
    written = 0
    uploader = str(uploaded_by or "").strip()[:128]
    for raw in rows:
        if not isinstance(raw, dict):
            raise ValueError("each item must be an object")
        chapter_id = str(raw.get("chapter_id") or "").strip()
        question_type = str(raw.get("question_type") or "").strip()
        question = str(raw.get("question") or "").strip()
        if not chapter_id or question_type not in VALID_TYPES or not question:
            raise ValueError(
                "each item requires chapter_id, a valid question_type, and question"
            )
        fingerprint = _fingerprint_for(raw)
        options = raw.get("options_or_pairs", {})
        citations = raw.get("citations_json", raw.get("citations", []))
        existing = (
            TeacherQuestionBankItem.query.filter_by(
                chapter_id=chapter_id,
                question_type=question_type,
                question_fingerprint=fingerprint,
            )
            .order_by(TeacherQuestionBankItem.created_at.desc())
            .first()
        )
        if existing is None:
            existing = TeacherQuestionBankItem(
                id=uuid.uuid4().hex,
                chapter_id=chapter_id,
                question_type=question_type,
                question_fingerprint=fingerprint,
            )
            db.session.add(existing)

        existing.chapter = str(raw.get("chapter") or existing.chapter or "")[:256]
        existing.topic = str(raw.get("topic") or "")[:256]
        existing.question = question
        existing.options_or_pairs = _dumps(options)
        existing.correct_answer = str(raw.get("correct_answer") or "")
        existing.explanation = str(raw.get("explanation") or "")
        existing.citations_json = _dumps(citations if citations is not None else [])
        existing.source_document = str(raw.get("source_document") or "")[:255]
        page = raw.get("source_page")
        existing.source_page = int(page) if page is not None and str(page).strip() != "" else None
        existing.source_excerpt = str(raw.get("source_excerpt") or "")
        existing.source_chunk_id = str(raw.get("source_chunk_id") or "")[:160]
        existing.knowledge_key = str(raw.get("knowledge_key") or "")[:160]
        existing.fact_key = str(raw.get("fact_key") or "")[:80]
        existing.fact_summary = str(raw.get("fact_summary") or "")[:500]
        if uploader:
            existing.created_by = uploader
        elif not existing.created_by:
            existing.created_by = ""
        if "is_active" in raw:
            existing.is_active = bool(raw.get("is_active"))
        else:
            existing.is_active = True if existing.is_active is None else bool(existing.is_active)

        if quiz_payload_from_item(existing) is None:
            raise ValueError(
                f"item fingerprint={fingerprint} is missing type-specific answer fields in options_or_pairs"
            )
        written += 1

    db.session.commit()
    return written
