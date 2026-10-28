"""Behavioral tests for the structured teacher upload workflow."""
import io
import re
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect
from common.models import BankQuestion, BankQuestionDetails, db
from common.question_bank import parse_workbook
from common.question_bank_extended import QUESTION_HEADERS, DETAIL_HEADERS
from validation_tests.test_question_bank import client, upload


def extended(kind='MATCHING', question=None, details=None, reference=None):
    examples = {
        'MATCHING': ('Match the protocols.', '', 'Matching', [['Q1','L1','HTTP','R1','Web content'],['Q1','L2','DNS','R2','Domain resolution']]),
        'FILL': ('DNS resolves names to {{blank1}} addresses on port {{blank2}}.', '', 'Blanks', [['Q1','blank1','IP',True],['Q1','blank1','Internet Protocol',True],['Q1','blank2',53,False]]),
        'SHORT': ('Explain two TCP/UDP differences.', 'TCP is connection-oriented and reliable; UDP is connectionless with no delivery guarantees.', 'Rubric', [['Q1','C1','Connection-oriented versus connectionless','1.5'],['Q1','C2','Reliable delivery versus no guarantee','0.5']]),
    }
    prompt, ref, sheet_name, rows = examples[kind]
    book=Workbook();book.active.title='Questions';book.active.append(QUESTION_HEADERS)
    book.active.append(['Q1',kind,'Networks',prompt if question is None else question,ref if reference is None else reference,'Teacher explanation'])
    sheet=book.create_sheet(sheet_name);sheet.append(DETAIL_HEADERS[sheet_name])
    for row in rows if details is None else details:sheet.append(row)
    output=io.BytesIO();book.save(output);book.close();return output.getvalue()


def edit_book(content, change):
    book=load_workbook(io.BytesIO(content));change(book)
    output=io.BytesIO();book.save(output);book.close();return output.getvalue()


def get_token(response):
    match=re.search('name="preview_token" value="([^"]+)"',response.text)
    return match.group(1) if match else ''


@pytest.mark.parametrize('kind',['MATCHING','FILL','SHORT'])
def test_template_preview_confirm_and_list(client,kind):
    c,app=client
    response=c.get('/question-bank/template?type='+kind)
    questions,errors=parse_workbook(response.data)
    assert not errors and questions[0]['question_type']==kind
    # Use independent authored content as well as the downloaded example.
    content=extended(kind)
    response=upload(c,'preview',content=content)
    token=get_token(response)
    assert response.status_code==200 and token
    with app.app_context():assert BankQuestion.query.count()==0 and BankQuestionDetails.query.count()==0
    response=upload(c,'confirm',content=content,token=token)
    assert response.status_code==200 and response.json['count']==1
    with app.app_context():
        question=BankQuestion.query.one();data=question.details.data
        assert question.question_type==kind and question.source=='teacher_upload'
        assert question.correct_answer=='' and question.options=={}
        assert data['schema_version']==1
        if kind=='MATCHING':assert data['correct_pairs']=={'L1':'R1','L2':'R2'}
        elif kind=='FILL':
            assert data['blanks'][0]['accepted_answers']==['IP','Internet Protocol']
            assert data['blanks'][0]['ignore_case'] is True
            assert data['blanks'][1]['accepted_answers']==['53']
        else:
            assert data['max_score']=='2.0' and len(data['rubric'])==2
            assert 'reference_answer' in data
    page=c.get('/question-bank/').text
    assert 'Q1' in page and 'Teacher explanation' in page
    if kind=='MATCHING':assert 'L1 ↔ R1' in page
    elif kind=='FILL':assert 'Internet Protocol' in page and 'Case-sensitive' in page
    else:assert '2.0 points total' in page and 'Reference answer' in page
    assert not get_token(upload(c,'preview',content=content))


@pytest.mark.parametrize('details,expected',[
    ([['Q1','L1','HTTP','R1','Web']], '2–20'),
    ([['Q1','L1','HTTP','R1','Web'],['Q1','L1','DNS','R2','Names']], 'Duplicate left_id'),
    ([['Q1','L1','HTTP','R1','Web'],['Q1','L2','DNS','R1','Names']], 'Duplicate right_id'),
    ([['Q1','L1','HTTP','R1','Web'],['Q1','L2','DNS','R2','web']], 'Duplicate right_text'),
    ([['Q1','L1','HTTP','R1','Web'],['Q1','L2','','R2','Names']], 'both texts'),
])
def test_matching_rejects_ambiguous_or_incomplete_pairs(details,expected):
    _,errors=parse_workbook(extended(details=details))
    assert expected in str(errors)
    assert all('sheet' in e and 'row' in e for e in errors)


@pytest.mark.parametrize('question,details,expected',[
    ('No blank here.',None,'placeholders'),
    ('{{blank1}} and {{blank1}}',None,'unique'),
    ('{{blank0}}',None,'Invalid placeholder'),
    ('{{blank1}} and {{blank2}}',[['Q1','blank1','IP',True]],'Missing accepted answer for blank2'),
    ('{{blank1}}',[['Q1','blank2','IP',True]],'no matching placeholder'),
    ('{{blank1}}',[['Q1','blank1','',True]],'accepted_answer is required'),
    ('{{blank1}}',[['Q1','blank1','IP','maybe']],'ignore_case must'),
    ('{{blank1}}',[['Q1','blank1','IP',True],['Q1','blank1','ip',True]],'Duplicate accepted'),
    ('{{blank1}}',[['Q1','blank1','IP',True],['Q1','blank1','Internet Protocol',False]],'same ignore_case'),
])
def test_fill_validates_each_placeholder_and_answer(question,details,expected):
    _,errors=parse_workbook(extended('FILL',question=question,details=details))
    assert expected in str(errors)


def test_case_sensitive_variants_are_distinct():
    questions,errors=parse_workbook(extended('FILL',question='{{blank1}}',details=[['Q1','blank1','IP',False],['Q1','blank1','ip',False]]))
    assert not errors and len(questions[0]['details']['blanks'][0]['accepted_answers'])==2


@pytest.mark.parametrize('points',['0','-1','NaN','Infinity','1.001','1e100000','1001','abc','=1+1'])
def test_short_rubric_rejects_invalid_points(points):
    _,errors=parse_workbook(extended('SHORT',details=[['Q1','C1','A criterion',points]]))
    assert errors


def test_short_requires_reference_and_unique_rubric():
    _,errors=parse_workbook(extended('SHORT',reference='',details=[]))
    assert 'reference_answer' in str(errors) and '1–20' in str(errors)
    _,errors=parse_workbook(extended('SHORT',details=[['Q1','C1','Point one',600],['Q1','c1','',600]]))
    assert 'unique' in str(errors) and 'criterion is required' in str(errors) and '1000' in str(errors)


def test_missing_sheet_and_orphan_rows():
    data=edit_book(extended(),lambda book:book.remove(book['Matching']))
    assert 'Missing Matching sheet' in str(parse_workbook(data)[1])
    data=extended(details=[['missing','L1','HTTP','R1','Web'],['Q1','L2','DNS','R2','Names']])
    assert 'question_id must refer' in str(parse_workbook(data)[1])


def test_formula_in_details_blocks_entire_import(client):
    c,app=client
    content=extended(details=[['Q1','L1','=1+1','R1','Web'],['Q1','L2','DNS','R2','Names']])
    response=upload(c,'preview',content=content)
    assert 'Matching · Excel row 2' in response.text and 'Formulas' in response.text
    assert not get_token(response)
    assert upload(c,'confirm',content=content).status_code==400
    with app.app_context():assert BankQuestion.query.count()==0 and BankQuestionDetails.query.count()==0


def test_active_detail_sheet_and_invalid_headers():
    content=edit_book(extended(),lambda b:setattr(b,'active',1))
    assert not parse_workbook(content)[1]
    content=edit_book(extended(),lambda b:setattr(b['Matching']['A1'],'value','wrong'))
    with pytest.raises(ValueError,match='Matching'):parse_workbook(content)


def test_structured_confirmation_bound_to_answers(client):
    c,app=client;content=extended('SHORT')
    token=get_token(upload(c,'preview',content=content))
    changed=edit_book(content,lambda b:setattr(b['Rubric']['D2'],'value',10))
    assert upload(c,'confirm',content=changed,token=token).status_code==400
    with app.app_context():assert BankQuestionDetails.query.count()==0


def test_additive_details_table_preserves_existing_bank(client):
    c,app=client
    with app.app_context():
        db.session.add(BankQuestion(owner_id=1,course='OLD',question_id='TF-old',question_type='TF',question='Existing question',options={},correct_answer='TRUE'))
        db.session.commit()
        BankQuestionDetails.__table__.drop(db.engine)
        db.create_all()
        assert 'teacher_bank_question_details' in inspect(db.engine).get_table_names()
        assert BankQuestion.query.filter_by(question_id='TF-old').one().correct_answer=='TRUE'
    content=extended('SHORT');token=get_token(upload(c,'preview',content=content))
    assert upload(c,'confirm',content=content,token=token).status_code==200


def test_structured_answers_remain_teacher_only_and_escaped(client):
    c,_=client
    content=extended('SHORT',reference='<script>alert(1)</script>')
    response=upload(c,'preview',content=content)
    assert '&lt;script&gt;' in response.text and '<script>' not in response.text
    token=get_token(response);assert upload(c,'confirm',content=content,token=token).status_code==200
    with c.session_transaction() as s:s['_user_id']='2'
    assert 'alert(1)' not in c.get('/question-bank/').text
    with c.session_transaction() as s:s['_user_id']='3'
    assert c.get('/question-bank/').status_code==403


def test_extended_template_links_honor_proxy(client):
    c,_=client;response=c.get('/question-bank/',headers={'X-Forwarded-Prefix':'/teacher'})
    assert '/teacher/question-bank/template?type=MATCHING' in response.text
    assert '/teacher/question-bank/template?type=FILL' in response.text
    assert '/teacher/question-bank/template?type=SHORT' in response.text
    assert c.get('/question-bank/template?type=UNKNOWN').status_code==400


def test_all_three_types_can_share_one_structured_workbook(client):
    c,app=client;book=Workbook();book.active.title='Questions';book.active.append(QUESTION_HEADERS)
    for kind in ('MATCHING','FILL','SHORT'):
        source=load_workbook(io.BytesIO(extended(kind)))
        main=list(source['Questions'].values)[1];main=list(main);main[0]=kind+'1';book.active.append(main)
        for sheet in source.worksheets[1:]:
            target=book.create_sheet(sheet.title)
            for index,row in enumerate(sheet.values):
                values=list(row)
                if index:values[0]=kind+'1'
                target.append(values)
        source.close()
    output=io.BytesIO();book.save(output);book.close();content=output.getvalue()
    token=get_token(upload(c,'preview',content=content));assert token
    assert upload(c,'confirm',content=content,token=token).json['count']==3
    with app.app_context():assert BankQuestion.query.count()==3 and BankQuestionDetails.query.count()==3


def test_broken_sheet_returns_validation_error(client):
    import zipfile
    c,_=client;content=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(extended())) as original, zipfile.ZipFile(content,'w') as broken:
        for name in original.namelist():
            broken.writestr(name,b'<invalid' if name=='xl/worksheets/sheet2.xml' else original.read(name))
    response=upload(c,'preview',content=content.getvalue())
    assert response.status_code==400 and 'Invalid' in response.json['error'] and 'workbook' in response.json['error']
